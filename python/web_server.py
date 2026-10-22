#!/usr/bin/env python3
"""PDP Digital Route - local web app.

Serves a premium dashboard and runs the route solver on demand.
"""
import base64
import json
import logging
import os
import sys
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path


def _load_settings():
    try:
        s = json.loads((Path(__file__).resolve().parent / "settings.json").read_text(encoding="utf-8"))
        return s
    except Exception:
        return {
            "warehouse_name": "Warehouse",
            "warehouse_lat": 0.0,
            "warehouse_lon": 0.0,
            "pdp_value_cr": 10,
            "pdp_volume_mt": 12,
            "truck_capacity_mt": 12,
            "brand": "Syngenta",
            "title": "PDP Digital Route",
        }

SETTINGS = _load_settings()

if getattr(sys, "frozen", False):
    BASE = Path(sys._MEIPASS)
    WORK = Path(sys.executable).resolve().parent
else:
    BASE = Path(__file__).resolve().parent
    WORK = BASE
ROOT = BASE
UPLOAD = WORK / "uploads"
OUT = WORK / "data"
UPLOAD.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)
os.chdir(str(WORK))
sys.path.insert(0, str(ROOT))

LOG_DIR = WORK / "logs"
LOG_DIR.mkdir(exist_ok=True)
logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s %(levelname)s %(message)s",
                    handlers=[logging.FileHandler(LOG_DIR / "server.log", encoding="utf-8")])

# import solver modules
import candidates
import solver
import route_opt
from build_data import build

LOCK = threading.Lock()
STATE = {"status": "idle", "message": ""}


def run_solve(gps_path, orders_path, api_key, cap, vmin, pmin):
    """Run the full pipeline. Returns (ok, result_json)."""
    data = build(gps_path, orders_path, out_dir=str(OUT), api_key=api_key or "")
    cands = candidates.generate_candidates(data, cap_max=cap, v_min=vmin, p_min=pmin)
    result = solver.solve(data, cands, cap_max=cap, v_min=vmin, p_min=pmin)
    # SA stop-order optimization on road matrix
    if result.get("status") == "Optimal":
        dm = data["distMatrix"]
        for route in result["routes"]:
            order = [c["id"] for c in route["customers"]]
            orig_dist = candidates.route_distance(order, dm)
            new_order = route_opt.optimize_order(order, dm, iterations=4000)
            m = {c["id"]: c for c in route["customers"]}
            route["orig_order"] = order
            route["orig_distance"] = orig_dist
            route["customers"] = [m[x] for x in new_order]
            route["distance"] = candidates.route_distance(new_order, dm)
        result["totalDistance"] = round(sum(r["distance"] for r in result["routes"]), 2)
    # google maps urls
    wh = data["warehouse"]
    cust_lookup = {c["id"]: c for c in data["customers"]}
    for r in result["routes"]:
        parts = [f"{wh['lat']},{wh['lon']}"]
        for c in r["customers"]:
            cc = cust_lookup.get(c["id"], {})
            if "lat" in cc and "lon" in cc:
                parts.append(f"{cc['lat']},{cc['lon']}")
        r["gmaps_url"] = "https://www.google.com/maps/dir/" + "/".join(parts)
    return True, {"data": data, "result": result}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        logging.info("%s - %s", self.address_string(), fmt % args)

    def _send(self, code, body, ctype="application/json"):
        if isinstance(body, (dict, list)):
            body = json.dumps(body).encode()
        elif isinstance(body, str):
            body = body.encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _file(self, path, ctype):
        p = ROOT / path
        if not p.exists():
            return self._send(404, "not found", "text/plain")
        self._send(200, p.read_bytes(), ctype)

    def do_GET(self):
        path = self.path.split("?")[0]
        if path == "/":
            return self._file("index.html", "text/html; charset=utf-8")
        if path.startswith("/assets/"):
            name = path[len("/assets/"):]
            p = ROOT / "assets" / name
            if not p.exists():
                return self._send(404, "not found", "text/plain")
            ctype = {"html": "text/html", "js": "application/javascript",
                     "css": "text/css", "svg": "image/svg+xml",
                     "woff2": "font/woff2", "json": "application/json"}.get(
                         p.suffix.lstrip("."), "application/octet-stream")
            return self._file("assets/" + name, ctype)
        if path == "/api/state":
            return self._send(200, STATE)
        if path == "/api/settings":
            return self._send(200, SETTINGS)
        return self._send(404, "not found", "text/plain")

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        raw = self.rfile.read(length) if length else b""
        try:
            body = json.loads(raw or b"{}")
        except Exception:
            return self._send(400, {"error": "bad json"})

        if self.path == "/api/upload":
            name = body.get("name", "file")
            b64 = body.get("content", "")
            try:
                data = base64.b64decode(b64)
            except Exception:
                return self._send(400, {"error": "bad content"})
            (UPLOAD / name).write_bytes(data)
            return self._send(200, {"ok": True, "path": str(UPLOAD / name)})

        if self.path == "/api/export":
            try:
                import openpyxl
                data = json.loads(body.get("data", "{}"))
                result = json.loads(body.get("result", "{}"))
                wb = openpyxl.Workbook()
                cust_lookup = {c["id"]: c for c in data["customers"]}
                ws = wb.active; ws.title = "Route Summary"
                ws.append(["Route ID", "Customer Count", "Total Volume (MT)", "Total Value (BDT)",
                           "Road Distance (km)", "PDP Status", "Capacity Utilisation"])
                for i, r in enumerate(result["routes"]):
                    ws.append([f"R{i+1:02d}", len(r["customers"]), r["volume"], r["value"],
                               r["distance"], r.get("qualifiesBy", ""), r.get("utilisation", 0)])
                ws2 = wb.create_sheet("Route Stops")
                ws2.append(["Route ID", "Stop", "BP ID", "Customer Name", "Volume (MT)", "Value (BDT)"])
                for i, r in enumerate(result["routes"]):
                    for j, c in enumerate(r["customers"]):
                        cc = cust_lookup.get(c["id"], {})
                        ws2.append([f"R{i+1:02d}", j + 1, cc.get("bpId", ""), cc.get("name", ""),
                                    cc.get("orderVolume", 0), cc.get("orderValue", 0)])
                ws3 = wb.create_sheet("Google Maps Links")
                ws3.append(["Route ID", "URL"])
                for r in result["routes"]:
                    ws3.append([r.get("rid", ""), r.get("gmaps_url", "")])
                import io
                buf = io.BytesIO(); wb.save(buf); buf.seek(0)
                raw = buf.read()
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", 'attachment; filename="route-plan.xlsx"')
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
                return
            except Exception as e:
                logging.exception("export failed")
                return self._send(500, {"error": str(e)})

        if self.path == "/api/solve":
            gps = body.get("gps", "")
            orders = body.get("orders", "")
            api_key = body.get("api_key", "")
            cap = float(body.get("cap", 12))
            vmin = float(body.get("vmin", 12))
            pmin = float(body.get("pmin", 10)) * 10_000_000
            if not gps or not orders:
                return self._send(400, {"error": "missing files"})
            try:
                ok, payload = run_solve(gps, orders, api_key, cap, vmin, pmin)
                return self._send(200, {"ok": True, **payload})
            except Exception as e:
                logging.exception("solve failed")
                return self._send(500, {"error": str(e)})

        return self._send(404, {"error": "unknown"})


def main():
    import webbrowser
    import socket
    port = 8765
    srv = None
    for p in (8765, 8766, 8767, 8768, 8769):
        try:
            srv = ThreadingHTTPServer(("127.0.0.1", p), Handler)
            port = p
            break
        except OSError:
            continue
    if srv is None:
        logging.error("no free port")
        return
    url = f"http://127.0.0.1:{port}/"
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    logging.info("serving on %s", url)
    # wait until the socket accepts connections, then open the browser
    def open_when_ready():
        for _ in range(50):
            try:
                with socket.create_connection(("127.0.0.1", port), timeout=0.5):
                    webbrowser.open(url)
                    return
            except OSError:
                time.sleep(0.1)
    threading.Thread(target=open_when_ready, daemon=True).start()
    try:
        while True:
            time.sleep(3600)
    except KeyboardInterrupt:
        srv.shutdown()


if __name__ == "__main__":
    main()
